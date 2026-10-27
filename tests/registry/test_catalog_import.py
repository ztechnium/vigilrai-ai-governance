"""Catalog workbook import."""

from __future__ import annotations

import uuid
from datetime import datetime, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from rai_gateway.auth.passwords import hash_password
from rai_gateway.core.database import get_db
from rai_gateway.governance.domain import GovernanceRole, MembershipStatus
from rai_gateway.main import app
from rai_gateway.models.database import AdminCredential, UserAccount
from rai_gateway.models.governance import OrganizationMembership
from rai_gateway.models.registry import (
    ExternalService,
    MCPServer,
    ModelResource,
    ToolResource,
)
from rai_gateway.registry.catalog_import import SHEET_HEADERS
from rai_gateway.services.admin_console import _ensure_default_rbac_seed


@pytest.fixture()
def public_client(db_session):
    def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    try:
        yield client
    finally:
        app.dependency_overrides.pop(get_db, None)


@pytest.fixture()
def auth_headers(public_client, db_session) -> dict[str, str]:
    _ensure_default_rbac_seed(db_session)
    suffix = uuid.uuid4().hex[:8]
    email = f"catalog-{suffix}@example.com"
    user_id = f"user-{suffix}"
    db_session.add(
        UserAccount(
            user_id=user_id,
            email=email,
            name="Catalog User",
            status="active",
            email_verified_at=datetime.now(timezone.utc),
        )
    )
    db_session.flush()
    db_session.add(
        AdminCredential(
            email=email,
            user_id=user_id,
            password_hash=hash_password("SecurePass123!"),
            is_active=True,
        )
    )
    db_session.add(
        OrganizationMembership(
            user_id=user_id,
            role=GovernanceRole.ORGANIZATION_ADMIN.value,
            status=MembershipStatus.ACTIVE.value,
            created_by=user_id,
        )
    )
    db_session.commit()
    login = public_client.post(
        "/v1/auth/login",
        json={"email": email, "password": "SecurePass123!"},
    )
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def _workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catalog"
    sheet.append(list(SHEET_HEADERS))
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_catalog_template_and_import(public_client, db_session, auth_headers):
    template = public_client.get(
        "/v1/registry/catalog/import-template",
        headers=auth_headers,
    )
    assert template.status_code == 200
    book = load_workbook(BytesIO(template.content))
    assert book.sheetnames[:2] == ["Catalog", "Example"]
    sheet = book["Catalog"]
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert headers[:3] == ["type", "key", "name"]
    assert sheet.max_row == 1
    example = list(book["Example"].iter_rows(values_only=True))
    assert example[0][:3] == ("type", "key", "name")
    assert {row[0] for row in example[1:] if row[0]} == {
        "model",
        "mcp",
        "tool",
        "service",
    }

    suffix = uuid.uuid4().hex[:8]
    model_key = f"gpt-{suffix}"
    mcp_key = f"lab-mcp-{suffix}"
    tool_key = f"create-ticket-{suffix}"
    service_key = f"billing-{suffix}"
    content = _workbook(
        [
            [
                "model",
                model_key,
                "GPT",
                "openai",
                "Primary",
                "gpt-4.1",
                "2026-04",
                "",
                "",
                "",
                "",
                "",
            ],
            [
                "mcp",
                mcp_key,
                "Lab MCP",
                "lab",
                "Server",
                "",
                "",
                "http://example/mcp",
                "",
                "",
                "",
                "",
            ],
            [
                "tool",
                tool_key,
                "Create ticket",
                "lab",
                "Open",
                "",
                "",
                "",
                "support",
                "low",
                "no",
                mcp_key,
            ],
            [
                "service",
                service_key,
                "Billing",
                "internal",
                "API",
                "",
                "",
                "http://billing",
                "",
                "internal",
                "",
                "",
            ],
        ]
    )
    created = public_client.post(
        "/v1/registry/catalog/import",
        headers=auth_headers,
        files={
            "file": (
                "catalog.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert created.status_code == 200, created.text
    assert created.json() == {"created": 4, "updated": 0, "rows": 4}

    mcp = db_session.query(MCPServer).filter(MCPServer.server_key == mcp_key).one()
    tool = (
        db_session.query(ToolResource).filter(ToolResource.tool_key == tool_key).one()
    )
    assert tool.mcp_server_id == mcp.id
    assert (
        db_session.query(ModelResource)
        .filter(
            ModelResource.provider == "openai", ModelResource.model_name == "gpt-4.1"
        )
        .count()
        == 1
    )
    assert (
        db_session.query(ExternalService)
        .filter(ExternalService.service_key == service_key)
        .count()
        == 1
    )

    updated_book = _workbook(
        [
            [
                "model",
                model_key,
                "GPT",
                "openai",
                "Primary",
                "gpt-4.1",
                "2026-04",
                "",
                "",
                "",
                "",
                "",
            ],
            [
                "mcp",
                mcp_key,
                "Lab MCP",
                "lab",
                "Server",
                "",
                "",
                "http://example/mcp",
                "",
                "",
                "",
                "",
            ],
            [
                "tool",
                tool_key,
                "Create ticket v2",
                "lab",
                "Open",
                "",
                "",
                "",
                "support",
                "low",
                "yes",
                mcp_key,
            ],
            [
                "service",
                service_key,
                "Billing API",
                "internal",
                "API",
                "",
                "",
                "http://billing",
                "",
                "internal",
                "",
                "",
            ],
        ]
    )
    updated = public_client.post(
        "/v1/registry/catalog/import",
        headers=auth_headers,
        files={
            "file": (
                "catalog.xlsx",
                updated_book,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["created"] == 0
    assert updated.json()["updated"] == 4
    db_session.expire_all()
    tool = (
        db_session.query(ToolResource).filter(ToolResource.tool_key == tool_key).one()
    )
    assert tool.name == "Create ticket v2"
    assert tool.consequential_action is True
    service = (
        db_session.query(ExternalService)
        .filter(ExternalService.service_key == service_key)
        .one()
    )
    assert service.name == "Billing API"
